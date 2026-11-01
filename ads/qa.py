"""QA 多 Sheet 原始输入适配与必验集合核对；不维护第二套审批状态机。"""
import csv
import json
from pathlib import Path
import re

from .project import require, digest, canonical, path_in
from .evidence import load_current

ALIASES = {'用例编号': 'id', '编号': 'id', '用例名称': 'title', '标题': 'title',
           '前置条件': 'preconditions', '操作步骤': 'steps', '测试步骤': 'steps',
           '预期结果': 'expected', '优先级': 'priority', '模块': 'module'}


def rows(path):
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            content = list(csv.reader(stream))
        yield path.stem, content, set()
    elif path.suffix.lower() == '.xlsx':
        from zipfile import ZipFile
        with ZipFile(path) as archive:
            require(len(archive.infolist()) < 500 and sum(x.file_size for x in archive.infolist()) < 20*1024*1024,
                    'QA 工作簿超过读取边界')
            require(not any('vbaProject' in x or x.startswith('xl/externalLinks/') for x in archive.namelist()), '不接受宏或外部链接')
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError('读取 QA XLSX 需要可选依赖 openpyxl==3.1.5') from exc
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in book.worksheets:
                require(sheet.max_row <= 10000 and sheet.max_column <= 100, 'QA Sheet 超过读取边界')
                cells = list(sheet.iter_rows())
                formulas = {(r, c) for r, row in enumerate(cells, 1) for c, cell in enumerate(row)
                            if cell.data_type == 'f'}
                yield sheet.title, [[cell.value for cell in row] for row in cells], formulas
        finally:
            book.close()
    else:
        raise ValueError('QA 输入仅支持 CSV 或 XLSX；其他格式需要独立 Adapter')


def intake(path):
    path = Path(path)
    require(path.is_file() and not path.is_symlink() and path.stat().st_size <= 5*1024*1024, 'QA 输入缺失、链接或过大')
    cases, issues, sheets, ids = [], [], [], set()
    for sheet, content, formulas in rows(path):
        sheets.append(sheet)
        nonempty = [(i, row) for i, row in enumerate(content, 1) if any(x is not None and str(x).strip() for x in row)]
        if not nonempty:
            issues.append({'sheet': sheet, 'type': 'EMPTY_SHEET'})
            continue
        _, header = nonempty[0]
        names = [ALIASES.get(str(x or '').strip(), str(x or '').strip()) for x in header]
        require(len([x for x in names if x]) == len(set(x for x in names if x)), '重复 QA 列名')
        for number, row in nonempty[1:]:
            raw = {name: '' if i >= len(row) or row[i] is None else str(row[i]).strip() for i, name in enumerate(names) if name}
            module = raw.get('module') or sheet
            ident = module + '/' + (raw.get('id') or f'row-{number}')
            require(ident not in ids, 'QA 用例标识重复')
            ids.add(ident)
            defects = [name for name in ('title', 'steps', 'expected') if not raw.get(name)]
            if any(r == number for r, _ in formulas):
                defects.append('FORMULA_REQUIRES_REVIEW')
            case = {'id': ident, 'module': module, 'source_locator': f'{sheet}!R{number}',
                    'fields': raw, 'quality': 'VALID' if not defects else 'REVISION_REQUIRED',
                    'quality_issues': defects, 'case_sha256': digest(canonical(raw))}
            cases.append(case)
    require(bool(cases), '没有可审查 QA 用例')
    return {'schema': 'ads-qa-intake/v1', 'source_sha256': digest(path.read_bytes()), 'sheets': sheets,
            'cases': cases, 'issues': issues,
            'quality': 'VALID' if not issues and all(x['quality'] == 'VALID' for x in cases) else 'REVISION_REQUIRED',
            'requirement_reconciliation': 'NOT_REVIEWED', 'approval': 'NOT_GRANTED'}


def candidates(imported, contract):
    def tokens(text):
        words = set(re.findall(r'[a-z0-9_]{2,}', text.lower()))
        chars = ''.join(re.findall(r'[\u4e00-\u9fff]', text))
        return words | {chars[i:i+2] for i in range(max(0, len(chars)-1))}
    result = []
    for case in imported['cases']:
        source = tokens(' '.join(case['fields'].values()))
        matches = []
        for obligation in contract['obligations']:
            other = tokens(obligation.get('description', obligation['id']))
            score = len(source & other)
            if score:
                matches.append({'obligation': obligation['id'], 'overlap': score})
        result.append({'case_id': case['id'], 'mapping_status': 'CANDIDATE_ONLY',
                       'candidates': sorted(matches, key=lambda x: -x['overlap'])[:5]})
    return result


def reviewed_baseline(project):
    config = project.config.get('qa')
    require(isinstance(config, dict), '未配置独立 QA 输入')
    imported = intake(path_in(project.root, config['source'], exists=True))
    require(imported['quality'] == 'VALID', 'QA 质量审查存在未关闭问题')
    review = json.loads(path_in(project.root, config['review'], exists=True).read_text())
    require(review.get('schema') == 'ads-qa-review/v1' and review.get('source_sha256') == imported['source_sha256'], 'QA 审查未绑定当前原件')
    decisions = review.get('decisions')
    require(isinstance(decisions, list), '缺少外部 QA 决策投影')
    by_id = {row['case_id']: row for row in decisions}
    require(len(by_id) == len(decisions) and set(by_id) == {c['id'] for c in imported['cases']}, 'QA 决策遗漏或重复')
    for case in imported['cases']:
        decision = by_id[case['id']]
        require(decision.get('case_sha256') == case['case_sha256'] and bool(decision.get('reviewer'))
                and bool(decision.get('decision_ref')), 'QA 用例缺少版本绑定与独立决策引用')
        require(decision.get('disposition') in ('CONFIRMED', 'EXCLUDED'), 'QA 冲突或歧义尚未裁决')
        if decision['disposition'] == 'EXCLUDED':
            require(bool(decision.get('reason')), '排除 QA 用例必须保留明确理由')
        else:
            tests = decision.get('tests')
            require(isinstance(tests, list) and bool(tests), 'QA 用例缺少必验集合')
            identities = []
            for target in tests:
                require(isinstance(target, dict) and isinstance(target.get('check'), str) and target['check']
                        and isinstance(target.get('case'), str) and target['case'], 'QA 必验映射结构无效')
                identities.append((target['check'], target['case']))
            require(len(identities) == len(set(identities)), 'QA 必验集合包含重复测试映射')
    return imported, by_id


def evaluate(project):
    imported, by_id = reviewed_baseline(project)
    run = load_current(project, 'qa')
    cases = []
    for case in imported['cases']:
        decision = by_id[case['id']]
        require(decision.get('case_sha256') == case['case_sha256'] and bool(decision.get('reviewer'))
                and bool(decision.get('decision_ref')), 'QA 用例缺少版本绑定与独立决策引用')
        disposition = decision.get('disposition')
        require(disposition in ('CONFIRMED', 'EXCLUDED'), 'QA 冲突或歧义尚未裁决')
        if disposition == 'EXCLUDED':
            require(bool(decision.get('reason')), '排除 QA 用例必须保留明确理由')
            cases.append({'id': case['id'], 'status': 'EXCLUDED'})
            continue
        tests = decision.get('tests')
        require(isinstance(tests, list) and bool(tests), '已确认 QA 用例没有必验执行集合')
        outcomes, evidence_refs = [], []
        for target in tests:
            matches = [
                (check, observed)
                for check in run['checks'] if check['check_id'] == target['check']
                for observed in check['cases'] if observed['id'] == target['case']
            ]
            if len(matches) != 1:
                observed_status = 'MISSING' if not matches else 'AMBIGUOUS'
                passed, reference = False, None
            else:
                check, observed = matches[0]
                case_status = observed['status']
                if check['status'] != 'PASS' and case_status == 'PASS':
                    observed_status = 'CHECK_FAILED'
                else:
                    observed_status = case_status
                passed = check['status'] == 'PASS' and case_status == 'PASS'
                reference = {
                    'run_id': run['run_id'],
                    'subject_sha256': run['subject']['sha256'],
                    'check_id': target['check'],
                    'case_id': target['case'],
                    'check_status': check['status'],
                    'case_status': case_status,
                    'agent_evidence': observed.get('agent_evidence', {'status': 'NOT_AVAILABLE'}),
                    'report': check['report'],
                    'report_sha256': check['report_sha256'],
                }
            outcomes.append(passed)
            evidence_refs.append({
                'expected': 'PASS',
                'observed': observed_status,
                'status': 'PASS' if passed else 'NOT_VERIFIED',
                'reference': reference,
            })
        cases.append({
            'id': case['id'], 'module': case['module'],
            'source_locator': case['source_locator'],
            'case_sha256': case['case_sha256'],
            'expected': case['fields']['expected'],
            'decision_ref': decision['decision_ref'],
            'reviewer': decision['reviewer'],
            'status': 'PASS' if outcomes and all(outcomes) else 'NOT_VERIFIED',
            'verification_evidence': evidence_refs,
        })
    required = [case for case in cases if case['status'] != 'EXCLUDED']
    complete = bool(required) and run['status'] == 'PASS' and all(c['status'] == 'PASS' for c in required)
    return {'schema': 'ads-qa-execution/v1', 'case_execution': 'PASS_LOCAL' if complete else 'BLOCKED',
            'run_id': run['run_id'], 'subject_sha256': run['subject']['sha256'],
            'run_status': run['status'], 'cases': cases, 'source_sha256': imported['source_sha256'],
            'independence': 'NOT_VERIFIED', 'qa_acceptance': 'NOT_GRANTED',
            'note': '复用外部评审决定；本机执行不提供独立 QA 身份或验收授权'}


def defects(project):
    """只读导出失败/未执行差集，供既有缺陷流程消费，不新建缺陷状态。"""
    result = evaluate(project)
    imported, decisions = reviewed_baseline(project)
    run = load_current(project, 'qa')
    originals = {case['id']: case for case in imported['cases']}
    findings = []
    for case in result['cases']:
        if case['status'] in ('PASS', 'EXCLUDED'):
            continue
        original = originals[case['id']]
        observations = []
        for target in decisions[case['id']]['tests']:
            checks = [c for c in run['checks'] if c['check_id'] == target['check']]
            for check in checks:
                statuses = [c['status'] for c in check['cases'] if c['id'] == target['case']]
                observations.append({**target, 'check_status': check['status'],
                                     'test_statuses': statuses or ['MISSING'],
                                     'execution_failure': check['failure'],
                                     'report': check['report'], 'report_sha256': check['report_sha256']})
            if not checks:
                observations.append({**target, 'check_status': 'MISSING', 'test_statuses': ['NOT_RUN']})
        findings.append({'case_id': case['id'], 'source_locator': original['source_locator'],
                         'expected': original['fields']['expected'], 'observations': observations,
                         'route': 'bug-fix', 'cause': 'NOT_DIAGNOSED',
                         'next': '复现并区分产品缺陷、环境故障和规格冲突；保留独立预期'})
    return {'schema': 'ads-qa-defect-projection/v1', 'run_id': run['run_id'],
            'subject_sha256': run['subject']['sha256'], 'findings': findings,
            'status': 'BLOCKED' if findings else 'NO_EXECUTION_GAPS_LOCAL',
            'qa_acceptance': 'NOT_GRANTED', 'external_issues_created': False}
